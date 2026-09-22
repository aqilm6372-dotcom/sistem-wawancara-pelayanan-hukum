"""
Sistem Wawancara Pelayanan Hukum
---------------------------------
Aplikasi Flask lokal untuk:
1. Otentikasi/Login pengguna (HMAC) dengan proteksi sesi.
2. Menangkap cerita permasalahan hukum dari klien.
3. Merapikan cerita tersebut menjadi bahasa hukum formal menggunakan Gemini API.
4. Menampilkan pratinjau & fitur edit teks hasil AI di browser.
5. Menghasilkan dokumen Word (.docx) yang langsung diunduh dari RAM (BytesIO).
"""

import os
import hmac
import base64
import mimetypes
from functools import wraps
from datetime import datetime
from io import BytesIO

from flask import (
    Flask,
    render_template,
    request,
    send_file,
    jsonify,
    session,
    redirect,
    url_for,
)

from google import genai
from google.genai import types

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Initialize Flask App
app = Flask(__name__)

# Batas total ukuran file yang diupload per request (semua lampiran digabung: 20MB)
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024


@app.errorhandler(413)
def file_terlalu_besar(e):
    """Handler jika total upload file melebihi batas 20MB."""
    return jsonify(
        {"error": "Total ukuran file lampiran terlalu besar (maksimal 20MB)."}
    ), 413

# =========================================================================
# 1. KONFIGURASI LOGIN & PROTEKSI SESI
# =========================================================================
app.secret_key = os.environ.get(
    "FLASK_SECRET_KEY", "ganti-string-ini-dengan-string-acak-rahasia-anda"
)

LOGIN_USERNAME = os.environ.get("LOGIN_USERNAME", "kejaksaan")
LOGIN_PASSWORD = os.environ.get("LOGIN_PASSWORD", "kejarilangsa26")


def login_required(f):
    """Decorator untuk memastikan halaman/endpoint hanya bisa diakses setelah login."""
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not session.get("logged_in"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)

    return decorated_function


@app.route("/login", methods=["GET", "POST"])
def login():
    """Route untuk menampilkan dan memproses login pengguna."""
    error = None
    if request.method == "POST":
        username = request.form.get("username", "")
        password = request.form.get("password", "")
        
        # Pengecekan aman menggunakan HMAC compare_digest
        username_valid = hmac.compare_digest(username, LOGIN_USERNAME)
        password_valid = hmac.compare_digest(password, LOGIN_PASSWORD)
        
        if username_valid and password_valid:
            session["logged_in"] = True
            session["username"] = username
            return redirect(url_for("index"))
        
        error = "Username atau password salah."
    return render_template("login.html", error=error)


@app.route("/logout")
def logout():
    """Route untuk mengakhiri sesi login."""
    session.clear()
    return redirect(url_for("login"))


# =========================================================================
# 2. KONFIGURASI GEMINI API & PROMPT HUKUM
# =========================================================================
import os
from dotenv import load_dotenv

load_dotenv()

GEMINI_API_KEY = os.getenv("GEMINI_API_KEY")

client = genai.Client(api_key=GEMINI_API_KEY)
GEMINI_MODEL = "gemini-3.6-flash"

SYSTEM_PROMPT = """
Anda adalah asisten hukum profesional di Indonesia.
Tugas Anda: ubah cerita permasalahan berikut menjadi bahasa hukum formal
yang terstruktur, rapi, dan mudah dibaca.

Instruksi:
1. Perbaiki kosakata, ejaan, dan tata bahasa agar sesuai kaidah penulisan hukum formal.
2. Bagi tulisan menjadi beberapa paragraf yang runtut dan mudah dibaca
   (misalnya: paragraf pembuka/duduk perkara, kronologi kejadian, lalu
   permasalahan hukum yang timbul).
3. WAJIB pertahankan seluruh fakta asli yang disampaikan klien. Jangan
   menambah, mengurangi, mengarang, atau mengubah substansi fakta apa pun.
4. Hanya perbaiki CARA PENYAMPAIANNYA, bukan isi ceritanya.
5. Jika ada dokumen pendukung yang dilampirkan (gambar, PDF, atau file teks),
   baca dan manfaatkan informasi relevan di dalamnya untuk melengkapi
   kronologi (misalnya nomor sertifikat, tanggal, nama pihak). Tetap jangan
   mengarang fakta yang tidak didukung oleh cerita klien maupun lampiran.
6. Keluarkan HANYA teks hasil perbaikan dalam Bahasa Indonesia. Jangan
   tambahkan judul, salam pembuka, catatan, disclaimer, atau format
   markdown (tanpa tanda bintang, tanpa heading).
""".strip()


# =========================================================================
# 3. ROUTES APLIKASI
# =========================================================================
@app.route("/")
@login_required
def index():
    """Menampilkan halaman utama formulir wawancara."""
    return render_template("index.html")


@app.route("/generate", methods=["POST"])
@login_required
def generate():
    """Tahap 1: Memproses input cerita dan lampiran via Gemini AI.

    Mengembalikan data JSON untuk dipratinjau di browser.
    """
    kategori = request.form.get("kategori", "").strip()
    cerita = request.form.get("cerita", "").strip()
    file_list = request.files.getlist("dokumen")

    if not kategori or not cerita:
        return jsonify({"error": "Kategori dan cerita permasalahan wajib diisi."}), 400

    # Memproses lampiran yang diunggah
    try:
        teks_tambahan, media_parts, info_lampiran = proses_lampiran(file_list)
    except Exception as e:
        return jsonify({"error": f"Gagal memproses lampiran: {str(e)}"}), 500

    # Memanggil model Gemini AI
    try:
        teks_hukum = panggil_gemini(cerita, teks_tambahan, media_parts)
    except Exception as e:
        return jsonify({"error": f"Gagal memproses AI: {str(e)}"}), 500

    # Mengubah data lampiran (gambar) ke format base64 agar aman dikirim ke frontend JSON
    lampiran_summary = []
    for item in info_lampiran:
        data_item = {"nama": item["nama"], "jenis": item["jenis"]}
        if item.get("bytes"):
            data_item["bytes_b64"] = base64.b64encode(item["bytes"]).decode("utf-8")
        lampiran_summary.append(data_item)

    return jsonify({
        "status": "success",
        "kategori": kategori,
        "teks_hukum": teks_hukum,
        "lampiran": lampiran_summary,
    })


@app.route("/download_word", methods=["POST"])
@login_required
def download_word():
    """Tahap 2: Menerima teks laporan yang sudah ditinjau/diedit pengguna,

    kemudian mengembalikan file Word (.docx) dari memori RAM.
    """
    data = request.get_json() or {}
    kategori = data.get("kategori", "").strip()
    teks_hukum = data.get("teks_hukum", "").strip()
    lampiran_data = data.get("lampiran", [])

    if not teks_hukum:
        return jsonify({"error": "Teks laporan tidak boleh kosong."}), 400

    # Rekonstruksi array info_lampiran dari payload JSON
    info_lampiran = []
    for item in lampiran_data:
        bytes_data = None
        if item.get("bytes_b64"):
            try:
                bytes_data = base64.b64decode(item["bytes_b64"])
            except Exception:
                bytes_data = None
        info_lampiran.append({
            "nama": item.get("nama", ""),
            "jenis": item.get("jenis", ""),
            "bytes": bytes_data,
        })

    # Membuat file Word di dalam RAM
    try:
        file_stream = buat_dokumen_word(kategori, teks_hukum, info_lampiran)
    except Exception as e:
        return jsonify({"error": f"Gagal membuat dokumen: {str(e)}"}), 500

    nama_kategori_file = kategori.replace(" ", "_") if kategori else "Umum"
    nama_file = f"Laporan_Hukum_{nama_kategori_file}.docx"

    return send_file(
        file_stream,
        as_attachment=True,
        download_name=nama_file,
        mimetype=(
            "application/vnd.openxmlformats-officedocument"
            ".wordprocessingml.document"
        ),
    )


# =========================================================================
# 4. FUNGSI PEMPROSESAN DATA & DOKUMEN
# =========================================================================
EKSTENSI_GAMBAR = {".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp"}
EKSTENSI_PDF = {".pdf"}
EKSTENSI_WORD = {".docx"}
EKSTENSI_TEKS = {".txt"}


def proses_lampiran(file_list):
    """Membaca berbagai jenis file yang dilampirkan klien dan

    menyiapkannya untuk diproses oleh Gemini AI.
    """
    teks_tambahan_list = []
    media_parts = []
    info_lampiran = []

    for f in file_list:
        if not f or not f.filename:
            continue

        nama = f.filename
        ekstensi = os.path.splitext(nama)[1].lower()
        data = f.read()

        if not data:
            continue

        # File Gambar
        if ekstensi in EKSTENSI_GAMBAR:
            mime = mimetypes.guess_type(nama)[0] or "image/jpeg"
            media_parts.append(types.Part.from_bytes(data=data, mime_type=mime))
            info_lampiran.append({"nama": nama, "jenis": "Gambar", "bytes": data})

        # File PDF
        elif ekstensi in EKSTENSI_PDF:
            media_parts.append(
                types.Part.from_bytes(data=data, mime_type="application/pdf")
            )
            info_lampiran.append({"nama": nama, "jenis": "PDF", "bytes": None})

        # File Docx
        elif ekstensi in EKSTENSI_WORD:
            try:
                doc_dibaca = Document(BytesIO(data))
                teks_docx = "\n".join(
                    p.text for p in doc_dibaca.paragraphs if p.text.strip()
                )
                teks_tambahan_list.append(f"[Dari dokumen Word '{nama}']\n{teks_docx}")
                info_lampiran.append({"nama": nama, "jenis": "Word", "bytes": None})
            except Exception:
                info_lampiran.append(
                    {"nama": nama, "jenis": "Word (gagal dibaca)", "bytes": None}
                )

        # File Teks (.txt)
        elif ekstensi in EKSTENSI_TEKS:
            try:
                teks_txt = data.decode("utf-8", errors="ignore")
                teks_tambahan_list.append(f"[Dari file teks '{nama}']\n{teks_txt}")
                info_lampiran.append({"nama": nama, "jenis": "Teks", "bytes": None})
            except Exception:
                info_lampiran.append(
                    {"nama": nama, "jenis": "Teks (gagal dibaca)", "bytes": None}
                )

        # File format .doc lama
        elif ekstensi == ".doc":
            info_lampiran.append(
                {
                    "nama": nama,
                    "jenis": "Tidak didukung (Word 97-2003, simpan ulang sebagai .docx)",
                    "bytes": None,
                }
            )

        # Format lainnya
        else:
            info_lampiran.append(
                {"nama": nama, "jenis": "Tidak didukung (format tidak dikenali)", "bytes": None}
            )

    teks_tambahan = "\n\n".join(teks_tambahan_list)
    return teks_tambahan, media_parts, info_lampiran


def panggil_gemini(cerita: str, teks_tambahan: str = "", media_parts=None) -> str:
    """Mengirim permintaan analisis ke model Gemini API."""
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        temperature=0.4,
    )

    prompt_utama = cerita
    if teks_tambahan:
        prompt_utama += (
            "\n\n=== Isi Dokumen Pendukung yang Dilampirkan ===\n" + teks_tambahan
        )

    contents = [prompt_utama]
    if media_parts:
        contents.extend(media_parts)

    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=contents,
        config=config,
    )
    return response.text.strip()


def buat_dokumen_word(kategori: str, isi_teks: str, info_lampiran=None) -> BytesIO:
    """Menyusun dokumen Word secara dinamik di memori (BytesIO) dan

    menerapkan format perataan serta tata letak profesional.
    """
    doc = Document()

    # Judul Dokumen
    heading = doc.add_heading(f"Laporan Wawancara Hukum: {kategori}", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Timestamp Pembuatan
    tanggal = datetime.now().strftime("%d %B %Y, %H:%M WIB")
    p_tanggal = doc.add_paragraph()
    p_tanggal.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_tanggal = p_tanggal.add_run(f"Dicetak pada: {tanggal}")
    run_tanggal.italic = True
    run_tanggal.font.size = Pt(10)

    doc.add_paragraph()

    # Sub-Judul Uraian
    doc.add_heading("Uraian Permasalahan", level=2)

    # Isi Teks Hasil AI (Justify, Font Size 12)
    paragraf_list = [p.strip() for p in isi_teks.split("\n") if p.strip()]
    for paragraf in paragraf_list:
        p = doc.add_paragraph(paragraf)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.size = Pt(12)

    # Sub-Judul Lampiran jika ada
    if info_lampiran:
        doc.add_paragraph()
        doc.add_heading("Lampiran Dokumen", level=2)

        for item in info_lampiran:
            p = doc.add_paragraph()
            run = p.add_run(f"• {item['nama']}  —  {item['jenis']}")
            run.font.size = Pt(11)

            # Sisipkan pratinjau gambar jika tipe lampiran adalah gambar
            if item["jenis"] == "Gambar" and item.get("bytes"):
                try:
                    doc.add_picture(BytesIO(item["bytes"]), width=Inches(4))
                except Exception:
                    doc.add_paragraph("   (Gagal menampilkan pratinjau gambar)")

    # Simpan ke stream RAM
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# =========================================================================
# RUN APPLICATION
# =========================================================================
if __name__ == "__main__":
    app.run(debug=True, host="127.0.0.1", port=5000, use_reloader=False)
